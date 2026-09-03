from __future__ import annotations

import io
import json
import os
import tempfile
import unittest
from datetime import datetime, timedelta, UTC
from email.utils import format_datetime
from pathlib import Path
from unittest.mock import MagicMock, patch

import openpyxl
from fastapi.testclient import TestClient

from preflight.api.app import app
from preflight.intake.email import (
    MAX_ATTACHMENT_BYTES,
    EmailAttachment,
    EmailIntakeConfig,
    EmailIntakeService,
    EmailMessageItem,
)
from preflight.models import CustomerMaster
from preflight.store import AuditStore


def create_sample_po_xlsx_bytes(po_number: str = "PO-2026-9999", sku: str = "LAPTOP-A14", qty: int = 2) -> bytes:
    """Helper creating valid PO Excel spreadsheet bytes."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DonHang"
    ws["A1"] = "ĐƠN HÀNG"
    ws["A2"] = "Mã PO:"
    ws["B2"] = po_number
    ws["A3"] = "Khách Hàng:"
    ws["B3"] = "Northstar Retail"
    ws["A4"] = "Tiền Tệ:"
    ws["B4"] = "VND"
    
    ws.append([])
    ws.append(["STT", "Mã SKU", "Tên Hàng", "Số Lượng", "ĐVT", "Đơn Giá", "Thành Tiền"])
    ws.append([1, sku, "Laptop A14 High Performance", qty, "PCS", 18500000, 18500000 * qty])
    
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class MockSMTPClient:
    def __init__(self):
        self.sent_messages = []

    def send_email(self, to_email: str, subject: str, body: str, attachment_path: str | Path | None = None) -> bool:
        self.sent_messages.append({
            "to": to_email,
            "subject": subject,
            "body": body,
            "attachment": str(attachment_path) if attachment_path else None,
        })
        return True


class MockIMAPClient:
    """Mailbox mock that honours \\Seen, so re-polling behaves like a real server."""

    def __init__(self, messages: list[EmailMessageItem] | None = None):
        self.messages = messages or []
        self.seen: set[str] = set()

    def fetch_messages(self) -> list[EmailMessageItem]:
        return [m for m in self.messages if m.raw_uid not in self.seen]

    def mark_seen(self, uids: list[str]) -> int:
        self.seen.update(uids)
        return len(uids)


class TestEmailIntake(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.db_path = Path(self.temp_dir.name) / "test_email_preflight.db"
        self.catalog_path = Path("examples/catalog.csv")
        self.store = AuditStore(self.db_path)
        self.smtp_client = MockSMTPClient()
        self.service = EmailIntakeService(
            store=self.store,
            smtp_client=self.smtp_client,
        )

    def tearDown(self):
        self.store.close()
        self.temp_dir.cleanup()

    def test_single_mail_two_xlsx_attachments_creates_two_orders_and_confirms(self):
        """1 email with 2 xlsx attachments -> 2 orders created, status PROCESSED, auto-reply confirmation sent."""
        xlsx_1 = create_sample_po_xlsx_bytes("PO-EMAIL-001", "LAPTOP-A14", 2)
        xlsx_2 = create_sample_po_xlsx_bytes("PO-EMAIL-002", "MOUSE-WL", 5)

        msg = EmailMessageItem(
            message_id="msg-uuid-1001@northstar.vn",
            sender_name="Purchasing Team",
            sender_email="orders@northstar.vn",
            subject="Đơn hàng tháng 9 - Đợt 1",
            date_str=datetime.now(UTC).isoformat(),
            attachments=[
                EmailAttachment(filename="PO_001.xlsx", content_bytes=xlsx_1),
                EmailAttachment(filename="PO_002.xlsx", content_bytes=xlsx_2),
            ],
        )

        result = self.service.process_message(msg, catalog_path=self.catalog_path)

        self.assertEqual(result.status, "PROCESSED")
        self.assertEqual(result.orders_created, 2)
        self.assertEqual(len(result.analysis_ids), 2)
        self.assertEqual(result.customer_resolved, "Northstar Retail")
        self.assertTrue(result.auto_reply_sent)

        # Check DB log
        log = self.store.get_email_inbox_log("msg-uuid-1001@northstar.vn")
        self.assertIsNotNone(log)
        self.assertEqual(log["status"], "PROCESSED")
        self.assertEqual(log["orders_created"], 2)

        # Check SMTP auto-reply sent
        self.assertEqual(len(self.smtp_client.sent_messages), 1)
        sent = self.smtp_client.sent_messages[0]
        self.assertEqual(sent["to"], "orders@northstar.vn")
        self.assertIn("Tiếp nhận PO thành công", sent["subject"])
        self.assertIn("PO-EMAIL-001", sent["body"])
        self.assertIn("PO-EMAIL-002", sent["body"])

    def test_idempotent_processing_duplicate_message_id_skipped(self):
        """Submitting the same message_id twice must be idempotent (no duplicate orders created)."""
        xlsx = create_sample_po_xlsx_bytes("PO-EMAIL-DUP", "LAPTOP-A14", 1)
        msg = EmailMessageItem(
            message_id="unique-msg-id-888@company.com",
            sender_name="Sender",
            sender_email="orders@northstar.vn",
            subject="PO Urgent",
            date_str=datetime.now(UTC).isoformat(),
            attachments=[EmailAttachment(filename="PO_Dup.xlsx", content_bytes=xlsx)],
        )

        res1 = self.service.process_message(msg, catalog_path=self.catalog_path)
        self.assertEqual(res1.status, "PROCESSED")
        self.assertEqual(res1.orders_created, 1)

        # Process same message again
        res2 = self.service.process_message(msg, catalog_path=self.catalog_path)
        self.assertEqual(res2.status, "DUPLICATE")
        self.assertEqual(len(self.smtp_client.sent_messages), 1)  # No second email sent

    def test_invalid_attachment_docx_ignored_and_auto_replies_with_template(self):
        """Email with only .docx attachment -> IGNORED, auto-reply sent with PO_MAU.xlsx template."""
        msg = EmailMessageItem(
            message_id="msg-docx-only@unknown.vn",
            sender_name="Khach Hang",
            sender_email="customer@unknown.vn",
            subject="Gửi đơn hàng bằng Word",
            date_str=datetime.now(UTC).isoformat(),
            attachments=[
                EmailAttachment(filename="DonHang.docx", content_bytes=b"FAKE DOCX BYTES"),
            ],
        )

        res = self.service.process_message(msg, catalog_path=self.catalog_path)
        self.assertEqual(res.status, "IGNORED")
        self.assertEqual(res.orders_created, 0)
        self.assertTrue(res.auto_reply_sent)

        # Check DB log
        log = self.store.get_email_inbox_log("msg-docx-only@unknown.vn")
        self.assertIsNotNone(log)
        self.assertEqual(log["status"], "IGNORED")

        # Check reply attached PO_MAU.xlsx
        self.assertEqual(len(self.smtp_client.sent_messages), 1)
        sent = self.smtp_client.sent_messages[0]
        self.assertEqual(sent["to"], "customer@unknown.vn")
        self.assertIn("Yêu cầu gửi lại đơn hàng", sent["subject"])
        self.assertIn("PO_MAU.xlsx", sent["attachment"])

    def test_customer_resolution_by_contact_emails(self):
        """Sender email in contact_emails maps to corresponding CustomerMaster."""
        # Add new customer with contact email
        cust = CustomerMaster(
            code="CUST-TEST-MAIL",
            name="Công Ty Phân Phối Ánh Dương",
            tax_code="0319998888",
            tier="VIP",
            aliases=["Ánh Dương", "Anh Duong Co"],
            contact_emails=["order@anhduong.vn", "procurement@anhduong.vn"],
        )
        self.store.create_customer(cust)

        found = self.store.get_customer_by_contact_email("procurement@anhduong.vn")
        self.assertIsNotNone(found)
        self.assertEqual(found.code, "CUST-TEST-MAIL")
        self.assertEqual(found.name, "Công Ty Phân Phối Ánh Dương")

        # Unknown email
        unknown = self.store.get_customer_by_contact_email("stranger@xyz.com")
        self.assertIsNone(unknown)

    def test_unresolved_customer_generates_finding(self):
        """Email from unregistered sender creates finding CUSTOMER_UNRESOLVED."""
        xlsx = create_sample_po_xlsx_bytes("PO-UNRESOLVED-01", "LAPTOP-A14", 1)
        msg = EmailMessageItem(
            message_id="msg-unresolved-123@stranger.org",
            sender_name="Stranger",
            sender_email="stranger@stranger.org",
            subject="Đơn hàng mới",
            date_str=datetime.now(UTC).isoformat(),
            attachments=[EmailAttachment(filename="Order.xlsx", content_bytes=xlsx)],
        )

        res = self.service.process_message(msg, catalog_path=self.catalog_path)
        self.assertEqual(res.status, "PROCESSED")
        self.assertEqual(res.orders_created, 1)

        # Inspect analysis in store
        analysis_id = res.analysis_ids[0]
        stored = self.store.get_order(analysis_id)
        self.assertIsNotNone(stored)
        findings = json.loads(stored["findings_json"])
        has_unresolved_finding = any(f.get("code") == "CUSTOMER_UNRESOLVED" for f in findings)
        self.assertTrue(has_unresolved_finding)

    def test_partial_failure_reports_bad_attachment_and_still_saves_good_one(self):
        """1 good + 1 corrupt xlsx -> PARTIAL, good order saved, bad file named in the reply."""
        good = create_sample_po_xlsx_bytes("PO-PARTIAL-OK", "LAPTOP-A14", 3)

        msg = EmailMessageItem(
            message_id="msg-partial-01@northstar.vn",
            sender_name="Purchasing Team",
            sender_email="orders@northstar.vn",
            subject="Đơn hàng có tệp lỗi",
            date_str=datetime.now(UTC).isoformat(),
            attachments=[
                EmailAttachment(filename="PO_Good.xlsx", content_bytes=good),
                EmailAttachment(filename="PO_Corrupt.xlsx", content_bytes=b"NOT A REAL XLSX"),
            ],
        )

        res = self.service.process_message(msg, catalog_path=self.catalog_path)

        # The readable attachment must still be ingested.
        self.assertEqual(res.status, "PARTIAL")
        self.assertEqual(res.orders_created, 1)
        self.assertEqual(res.po_numbers, ["PO-PARTIAL-OK"])

        # The unreadable one must be surfaced, not swallowed.
        self.assertEqual(len(res.failed_attachments), 1)
        self.assertEqual(res.failed_attachments[0][0], "PO_Corrupt.xlsx")

        log = self.store.get_email_inbox_log("msg-partial-01@northstar.vn")
        self.assertEqual(log["status"], "PARTIAL")
        self.assertEqual(log["orders_created"], 1)
        self.assertIsNotNone(log["error_message"])

        # The sender must be told which file failed.
        self.assertTrue(res.auto_reply_sent)
        sent = self.smtp_client.sent_messages[-1]
        self.assertIn("một phần", sent["subject"].lower())
        self.assertIn("PO_Corrupt.xlsx", sent["body"])
        self.assertIn("PO-PARTIAL-OK", sent["body"])
        self.assertIn("PO_MAU.xlsx", sent["attachment"])

    def test_partial_message_is_not_reprocessed_on_repoll(self):
        """A PARTIAL message must not be re-ingested, or its good orders duplicate."""
        good = create_sample_po_xlsx_bytes("PO-PARTIAL-DUP", "MOUSE-WL", 1)
        msg = EmailMessageItem(
            message_id="msg-partial-dup@northstar.vn",
            sender_name="Purchasing Team",
            sender_email="orders@northstar.vn",
            subject="Đơn hàng lặp một phần",
            date_str=datetime.now(UTC).isoformat(),
            attachments=[
                EmailAttachment(filename="Good.xlsx", content_bytes=good),
                EmailAttachment(filename="Bad.xlsx", content_bytes=b"BROKEN"),
            ],
        )

        first = self.service.process_message(msg, catalog_path=self.catalog_path)
        self.assertEqual(first.status, "PARTIAL")
        self.assertEqual(first.orders_created, 1)

        second = self.service.process_message(msg, catalog_path=self.catalog_path)
        self.assertEqual(second.status, "DUPLICATE")
        self.assertEqual(second.orders_created, 1)

    def test_all_attachments_unreadable_notifies_sender(self):
        """Every attachment corrupt -> ERROR, and the sender is told with the template."""
        msg = EmailMessageItem(
            message_id="msg-allbad-01@unknown.vn",
            sender_name="Khach Hang",
            sender_email="customer@unknown.vn",
            subject="Đơn hàng lỗi hoàn toàn",
            date_str=datetime.now(UTC).isoformat(),
            attachments=[
                EmailAttachment(filename="Bad1.xlsx", content_bytes=b"BROKEN ONE"),
                EmailAttachment(filename="Bad2.csv", content_bytes=b"\xff\xfe\x00garbage"),
            ],
        )

        res = self.service.process_message(msg, catalog_path=self.catalog_path)
        self.assertEqual(res.status, "ERROR")
        self.assertEqual(res.orders_created, 0)

        # Previously this path recorded a log but never told the customer.
        self.assertTrue(res.auto_reply_sent)
        sent = self.smtp_client.sent_messages[-1]
        self.assertEqual(sent["to"], "customer@unknown.vn")
        self.assertIn("Không tiếp nhận được đơn hàng", sent["subject"])
        self.assertIn("PO_MAU.xlsx", sent["attachment"])

        # An ERROR message is retryable, so a re-poll should try again.
        retry = self.service.process_message(msg, catalog_path=self.catalog_path)
        self.assertEqual(retry.status, "ERROR")

    def test_poll_once_caps_number_of_messages(self):
        """poll_once must not process an unbounded mailbox in one synchronous call."""
        def make(i: int) -> EmailMessageItem:
            return EmailMessageItem(
                message_id=f"msg-bulk-{i}@northstar.vn",
                sender_name="Purchasing Team",
                sender_email="orders@northstar.vn",
                subject=f"Đơn hàng {i}",
                date_str=datetime.now(UTC).isoformat(),
                attachments=[
                    EmailAttachment(
                        filename=f"PO_{i}.xlsx",
                        content_bytes=create_sample_po_xlsx_bytes(f"PO-BULK-{i}", "MOUSE-WL", 1),
                    )
                ],
            )

        service = EmailIntakeService(
            store=self.store,
            smtp_client=MockSMTPClient(),
            imap_client=MockIMAPClient([make(i) for i in range(10)]),
        )

        results = service.poll_once(catalog_path=self.catalog_path, max_messages=3)
        self.assertEqual(len(results), 3)

    def test_automated_sender_is_never_auto_replied(self):
        """A bounce must not be answered, or the reply loops back into this inbox."""
        for sender, label in [
            ("MAILER-DAEMON@googlemail.com", "bounce"),
            ("noreply@somebank.vn", "noreply"),
        ]:
            with self.subTest(sender=label):
                msg = EmailMessageItem(
                    message_id=f"auto-{label}@x.vn",
                    sender_name="Mail Delivery System",
                    sender_email=sender,
                    subject="Delivery Status Notification (Failure)",
                    date_str=datetime.now(UTC).isoformat(),
                    attachments=[],
                )
                res = self.service.process_message(msg, catalog_path=self.catalog_path)
                self.assertEqual(res.status, "IGNORED")
                self.assertFalse(res.auto_reply_sent)

        # Header-flagged auto-responders are caught even from a normal address.
        vacation = EmailMessageItem(
            message_id="auto-vacation@x.vn",
            sender_name="Nguoi Dung",
            sender_email="orders@northstar.vn",
            subject="Out of office",
            date_str=datetime.now(UTC).isoformat(),
            attachments=[],
            auto_submitted="auto-replied",
        )
        res = self.service.process_message(vacation, catalog_path=self.catalog_path)
        self.assertEqual(res.status, "IGNORED")
        self.assertFalse(res.auto_reply_sent)

        # Nothing at all should have been sent across every case above.
        self.assertEqual(self.smtp_client.sent_messages, [])

    def test_processed_messages_are_flagged_seen(self):
        """Handled mail must be flagged, or every poll re-fetches it forever."""
        good = create_sample_po_xlsx_bytes("PO-SEEN-001", "MOUSE-WL", 1)
        msg = EmailMessageItem(
            message_id="msg-seen-01@northstar.vn",
            sender_name="Purchasing Team",
            sender_email="orders@northstar.vn",
            subject="Đơn hàng cần đánh dấu",
            date_str=datetime.now(UTC).isoformat(),
            raw_uid="11",
            attachments=[EmailAttachment(filename="PO.xlsx", content_bytes=good)],
        )
        imap = MockIMAPClient([msg])
        service = EmailIntakeService(
            store=self.store, smtp_client=MockSMTPClient(), imap_client=imap
        )

        first = service.poll_once(catalog_path=self.catalog_path)
        self.assertEqual(len(first), 1)
        self.assertEqual(first[0].status, "PROCESSED")
        self.assertIn("11", imap.seen)

        # The mailbox no longer offers it, so the next poll has nothing to do.
        self.assertEqual(service.poll_once(catalog_path=self.catalog_path), [])

    def test_failed_message_is_retried_then_given_up(self):
        """A broken message retries a bounded number of times, then stops."""
        msg = EmailMessageItem(
            message_id="msg-retry-01@x.vn",
            sender_name="Khach",
            sender_email="customer@unknown.vn",
            subject="Tệp hỏng",
            date_str=datetime.now(UTC).isoformat(),
            raw_uid="21",
            attachments=[EmailAttachment(filename="bad.xlsx", content_bytes=b"BROKEN")],
        )
        imap = MockIMAPClient([msg])
        service = EmailIntakeService(
            store=self.store, smtp_client=MockSMTPClient(), imap_client=imap
        )

        # Failures stay unread so they can be retried.
        for _ in range(3):
            res = service.poll_once(catalog_path=self.catalog_path)
            self.assertEqual(res[0].status, "ERROR")
            self.assertNotIn("21", imap.seen)

        # Once the ceiling is hit the message is abandoned and flagged, so it
        # stops occupying the queue.
        res = service.poll_once(catalog_path=self.catalog_path)
        self.assertEqual(res[0].status, "GAVE_UP")
        self.assertIn("21", imap.seen)

    def test_new_order_is_not_starved_by_a_backlog_of_broken_mail(self):
        """A fresh PO must be processed even behind a large backlog of bad mail.

        Regression: the poll cap used to fill with unreadable mail every cycle,
        so a newly arrived order was never reached at all.
        """
        base = datetime.now(UTC) - timedelta(days=2)
        backlog = [
            EmailMessageItem(
                message_id=f"old-{i}@x.vn",
                sender_name="X",
                sender_email="someone@unknown.vn",
                subject=f"Tệp hỏng {i}",
                date_str=format_datetime(base + timedelta(minutes=i)),
                raw_uid=str(i),
                attachments=[EmailAttachment(filename=f"bad{i}.xlsx", content_bytes=b"BROKEN")],
            )
            for i in range(50)
        ]
        imap = MockIMAPClient(backlog)
        service = EmailIntakeService(
            store=self.store, smtp_client=MockSMTPClient(), imap_client=imap
        )
        service.poll_once(catalog_path=self.catalog_path)

        # The customer sends a real order into that backlog.
        fresh = EmailMessageItem(
            message_id="new-order@northstar.vn",
            sender_name="Purchasing Team",
            sender_email="orders@northstar.vn",
            subject="ĐƠN HÀNG MỚI",
            date_str=format_datetime(datetime.now(UTC)),
            raw_uid="999",
            attachments=[
                EmailAttachment(
                    filename="PO.xlsx",
                    content_bytes=create_sample_po_xlsx_bytes("PO-FRESH-001", "MOUSE-WL", 2),
                )
            ],
        )
        imap.messages.append(fresh)

        results = service.poll_once(catalog_path=self.catalog_path)
        processed = [r for r in results if r.message_id == "new-order@northstar.vn"]
        self.assertEqual(len(processed), 1, "đơn mới phải được xử lý ngay chu kỳ kế tiếp")
        self.assertEqual(processed[0].status, "PROCESSED")

    def test_backlog_eventually_drains(self):
        """Retries must still make progress, not be starved by fresh mail."""
        base = datetime.now(UTC) - timedelta(days=2)
        backlog = [
            EmailMessageItem(
                message_id=f"drain-{i}@x.vn",
                sender_name="X",
                sender_email="someone@unknown.vn",
                subject=f"Tệp hỏng {i}",
                date_str=format_datetime(base + timedelta(minutes=i)),
                raw_uid=str(i),
                attachments=[EmailAttachment(filename=f"bad{i}.xlsx", content_bytes=b"BROKEN")],
            )
            for i in range(30)
        ]
        imap = MockIMAPClient(backlog)
        service = EmailIntakeService(
            store=self.store, smtp_client=MockSMTPClient(), imap_client=imap
        )

        for _ in range(20):
            if not imap.fetch_messages():
                break
            service.poll_once(catalog_path=self.catalog_path)

        self.assertEqual(imap.fetch_messages(), [], "tồn đọng phải được dọn hết")

    def test_oversized_attachment_is_skipped(self):
        """A huge attachment must not be handed to the parser."""
        oversized = b"x" * (MAX_ATTACHMENT_BYTES + 1)
        raw = (
            b"From: Khach <customer@unknown.vn>\r\n"
            b"Subject: PO lon\r\n"
            b'Content-Type: multipart/mixed; boundary="B"\r\n\r\n'
            b"--B\r\n"
            b"Content-Type: application/octet-stream\r\n"
            b'Content-Disposition: attachment; filename="huge.xlsx"\r\n\r\n'
            + oversized
            + b"\r\n--B--\r\n"
        )
        parsed = self.service.parse_mime_message(raw)
        self.assertEqual(parsed.attachments, [])

    def test_dry_run_never_sends_mail_or_touches_the_mailbox(self):
        """Dry run must leave a real mailbox exactly as it was found."""
        msg = EmailMessageItem(
            message_id="dry-01@northstar.vn",
            sender_name="Purchasing Team",
            sender_email="orders@northstar.vn",
            subject="Đơn hàng thử nghiệm",
            date_str=datetime.now(UTC).isoformat(),
            raw_uid="31",
            attachments=[
                EmailAttachment(
                    filename="PO.xlsx",
                    content_bytes=create_sample_po_xlsx_bytes("PO-DRY-001", "MOUSE-WL", 1),
                )
            ],
        )
        imap = MockIMAPClient([msg])
        smtp = MockSMTPClient()
        service = EmailIntakeService(
            config=EmailIntakeConfig(dry_run=True),
            store=self.store,
            smtp_client=smtp,
            imap_client=imap,
        )

        results = service.poll_once(catalog_path=self.catalog_path)
        self.assertEqual(results[0].status, "PROCESSED")

        # No reply may leave the building, and nothing may be marked read.
        self.assertEqual(smtp.sent_messages, [])
        self.assertFalse(results[0].auto_reply_sent)
        self.assertEqual(imap.seen, set())

    def test_allowed_senders_ignores_unrelated_mail(self):
        """An allow-list makes a shared or personal inbox safe to poll."""
        wanted = EmailMessageItem(
            message_id="allow-01@northstar.vn",
            sender_name="Purchasing Team",
            sender_email="orders@northstar.vn",
            subject="Đơn hàng thật",
            date_str=datetime.now(UTC).isoformat(),
            raw_uid="41",
            attachments=[
                EmailAttachment(
                    filename="PO.xlsx",
                    content_bytes=create_sample_po_xlsx_bytes("PO-ALLOW-001", "MOUSE-WL", 1),
                )
            ],
        )
        # Personal mail that happens to sit in the same mailbox.
        unrelated = EmailMessageItem(
            message_id="allow-02@bank.vn",
            sender_name="Ngân hàng",
            sender_email="thongbao@nganhang.vn",
            subject="Sao kê tài khoản tháng 9",
            date_str=datetime.now(UTC).isoformat(),
            raw_uid="42",
            attachments=[EmailAttachment(filename="saoke.pdf", content_bytes=b"%PDF-1.4 fake")],
        )

        imap = MockIMAPClient([wanted, unrelated])
        smtp = MockSMTPClient()
        service = EmailIntakeService(
            config=EmailIntakeConfig(allowed_senders=frozenset({"orders@northstar.vn"})),
            store=self.store,
            smtp_client=smtp,
            imap_client=imap,
        )

        results = service.poll_once(catalog_path=self.catalog_path)
        by_id = {r.message_id: r for r in results}

        self.assertEqual(by_id["allow-01@northstar.vn"].status, "PROCESSED")
        self.assertEqual(by_id["allow-02@bank.vn"].status, "SKIPPED")

        # The bank must never be written to, and its mail must stay unread.
        self.assertEqual([m["to"] for m in smtp.sent_messages], ["orders@northstar.vn"])
        self.assertNotIn("42", imap.seen)

    def test_email_intake_api_routes(self):
        """Test GET /api/v1/intake/email/status and POST /api/v1/intake/email/poll."""
        client = TestClient(app, headers={"X-API-Key": "pf_dev_adm_9901"})

        # Test status endpoint
        res = client.get("/api/v1/intake/email/status")
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertIn("enabled", data)
        self.assertIn("imap_configured", data)
        self.assertIn("recent_error_count", data)

        # Test poll endpoint
        res_poll = client.post("/api/v1/intake/email/poll")
        self.assertEqual(res_poll.status_code, 200)
        poll_data = res_poll.json()
        self.assertEqual(poll_data["status"], "success")
        self.assertIn("processed_count", poll_data)

        # Test logs endpoint
        res_logs = client.get("/api/v1/intake/email/logs")
        self.assertEqual(res_logs.status_code, 200)
        self.assertIsInstance(res_logs.json(), list)


if __name__ == "__main__":
    unittest.main()
