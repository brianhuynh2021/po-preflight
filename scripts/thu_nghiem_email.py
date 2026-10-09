"""Try out the email PO intake flow right on your local machine.

No real IMAP/SMTP needed: the script builds a fake in-memory mailbox, simulates
the situations met in real life, then prints the result of each step.

    python scripts/thu_nghiem_email.py
"""
from __future__ import annotations

import io
import logging
import sys
from datetime import datetime, timedelta, UTC
from email.utils import format_datetime
from pathlib import Path

import openpyxl

from preflight.intake.email import (
    EmailAttachment,
    EmailIntakeService,
    EmailMessageItem,
)
from preflight.models import CustomerMaster
from preflight.store import AuditStore

logging.disable(logging.CRITICAL)  # keep output concise, only print the explanatory text

DB_PATH = Path("runtime/thu_nghiem_email.db")


def tao_file_po(po: str, khach: str, sku: str, so_luong: int, don_gia: int) -> bytes:
    """Create an Excel PO file that looks like one a real customer would send."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DonHang"
    ws["A1"] = "ĐƠN ĐẶT HÀNG"
    ws["A2"] = "Mã PO:"
    ws["B2"] = po
    ws["A3"] = "Khách Hàng:"
    ws["B3"] = khach
    ws["A4"] = "Tiền Tệ:"
    ws["B4"] = "VND"
    ws.append([])
    ws.append(["STT", "Mã SKU", "Tên Hàng", "Số Lượng", "ĐVT", "Đơn Giá", "Thành Tiền"])
    ws.append([1, sku, "Hàng hoá", so_luong, "PCS", don_gia, don_gia * so_luong])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class HopThuGia:
    """Fake mailbox: behaves like a real IMAP server, remembering read state."""

    def __init__(self, messages):
        self.messages = list(messages)
        self.seen: set[str] = set()

    def fetch_messages(self):
        return [m for m in self.messages if m.raw_uid not in self.seen]

    def mark_seen(self, uids):
        self.seen.update(uids)
        return len(uids)


class HopThuGuiDi:
    """Records mail the system sends to customers, standing in for real SMTP."""

    def __init__(self):
        self.sent = []

    def send_email(self, to_email, subject, body, attachment_path=None):
        self.sent.append(
            {"to": to_email, "subject": subject, "body": body, "attachment": attachment_path}
        )
        return True


def phan_cach(tieu_de: str) -> None:
    print("\n" + "=" * 72)
    print(f"  {tieu_de}")
    print("=" * 72)


def main() -> int:
    if DB_PATH.exists():
        DB_PATH.unlink()
    DB_PATH.parent.mkdir(exist_ok=True)

    store = AuditStore(DB_PATH)
    store.seed_default_customers()

    # Register the customer's email so the system can identify the sender automatically.
    store.create_customer(
        CustomerMaster(
            code="CUST-TEST",
            name="Công ty TNHH Thử Nghiệm",
            tax_code="0100000000",
            tier="STANDARD",
            contact_emails=["muahang@khachhang.vn"],
        )
    )

    now = datetime.now(UTC)
    thu = [
        # 1. Known customer sends a valid PO
        EmailMessageItem(
            message_id="thu-01@khachhang.vn",
            sender_name="Phòng Mua Hàng",
            sender_email="muahang@khachhang.vn",
            subject="Đơn đặt hàng tháng 9",
            date_str=format_datetime(now),
            raw_uid="1",
            attachments=[
                EmailAttachment(
                    filename="PO_T9.xlsx",
                    content_bytes=tao_file_po("PO-2026-001", "Công ty TNHH Thử Nghiệm", "LAPTOP-A14", 3, 18_500_000),
                )
            ],
        ),
        # 2. Customer sends a Word file - the system cannot read it
        EmailMessageItem(
            message_id="thu-02@khachhang.vn",
            sender_name="Phòng Mua Hàng",
            sender_email="muahang@khachhang.vn",
            subject="Đơn hàng gửi bằng Word",
            date_str=format_datetime(now - timedelta(minutes=5)),
            raw_uid="2",
            attachments=[EmailAttachment(filename="DonHang.docx", content_bytes=b"WORD")],
        ),
        # 3. Unknown sender, not registered in the customer master
        EmailMessageItem(
            message_id="thu-03@lachoac.vn",
            sender_name="Khách Lạ",
            sender_email="ai_do@lachoac.vn",
            subject="Đặt hàng",
            date_str=format_datetime(now - timedelta(minutes=10)),
            raw_uid="3",
            attachments=[
                EmailAttachment(
                    filename="PO_La.xlsx",
                    content_bytes=tao_file_po("PO-2026-002", "Công ty Lạ", "MOUSE-WL", 5, 450_000),
                )
            ],
        ),
        # 4. Automated bounce notice - must NEVER be replied to
        EmailMessageItem(
            message_id="thu-04@mailer",
            sender_name="Mail Delivery System",
            sender_email="MAILER-DAEMON@googlemail.com",
            subject="Delivery Status Notification (Failure)",
            date_str=format_datetime(now - timedelta(minutes=15)),
            raw_uid="4",
            attachments=[],
        ),
    ]

    hop_thu = HopThuGia(thu)
    smtp = HopThuGuiDi()
    service = EmailIntakeService(store=store, smtp_client=smtp, imap_client=hop_thu)

    phan_cach("LẦN QUÉT 1 — hộp thư có 4 thư chưa đọc")
    ket_qua = service.poll_once(catalog_path="examples/catalog.csv")

    dien_giai = {
        "PROCESSED": "✅ Đã tạo đơn thành công",
        "PARTIAL": "⚠️  Tạo đơn một phần, có tệp lỗi",
        "IGNORED": "⏭️  Bỏ qua (không có tệp hợp lệ hoặc là thư tự động)",
        "ERROR": "❌ Không đọc được tệp, sẽ thử lại",
        "GAVE_UP": "🛑 Đã thử nhiều lần, cần người kiểm tra",
        "DUPLICATE": "🔁 Thư đã xử lý trước đó",
    }
    for r in ket_qua:
        print(f"\n  Thư: {r.subject!r}")
        print(f"    Người gửi : {r.sender_email}")
        print(f"    Kết quả   : {dien_giai.get(r.status, r.status)}")
        if r.customer_resolved:
            print(f"    Khách hàng: {r.customer_resolved} (nhận diện qua email)")
        if r.po_numbers:
            print(f"    Mã PO     : {', '.join(r.po_numbers)}")
        print(f"    Trả lời KH: {'có' if r.auto_reply_sent else 'KHÔNG'}")

    phan_cach("MAIL HỆ THỐNG ĐÃ GỬI CHO KHÁCH")
    for m in smtp.sent:
        print(f"\n  → {m['to']}")
        print(f"    Tiêu đề: {m['subject']}")
        if m["attachment"]:
            print(f"    Đính kèm: {Path(str(m['attachment'])).name}")
    if not any(m["to"].lower().startswith("mailer-daemon") for m in smtp.sent):
        print("\n  ✅ Không trả lời thư tự động → không có nguy cơ lặp vô hạn")

    phan_cach("LẦN QUÉT 2 — kiểm tra không xử lý trùng")
    lan2 = service.poll_once(catalog_path="examples/catalog.csv")
    con_lai = len(hop_thu.fetch_messages())
    print(f"\n  Số thư xử lý lại : {len(lan2)}")
    print(f"  Còn chưa đọc     : {con_lai}")
    if not lan2:
        print("  ✅ Không thư nào bị xử lý lại — đúng như mong đợi")
    else:
        for r in lan2:
            print(f"    - {r.subject!r} → {dien_giai.get(r.status, r.status)}")
        print("  ℹ️  Thư lỗi được thử lại (đúng thiết kế), thư thành công thì không")

    phan_cach("ĐƠN HÀNG ĐÃ VÀO HỆ THỐNG")
    don = store.list_orders(limit=20)
    if not don:
        print("\n  (chưa có đơn nào)")
    for d in don:
        print(f"\n  • {d['po_number']} — {d['customer']}")
        print(f"    Trạng thái: {d['status']} | Giá trị: {d['total']}")
        chi_tiet = store.get_order(d["po_number"])
        import json
        for f in json.loads(chi_tiet["findings_json"]):
            print(f"      ⚠ [{f['code']}] {f['message'][:80]}")

    print(f"\n{'=' * 72}")
    print(f"  Dữ liệu thử nghiệm lưu tại: {DB_PATH}")
    print(f"  Xoá đi chạy lại: rm {DB_PATH}")
    print("=" * 72 + "\n")

    store.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
