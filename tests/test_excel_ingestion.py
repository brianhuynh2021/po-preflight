from __future__ import annotations

import io
import unittest
from decimal import Decimal

import openpyxl
from starlette.testclient import TestClient

from preflight.api.app import app
from preflight.ingestion.excel_parser import ExcelExtractor
from preflight.ingestion.pipeline import IntelligentIngestionPipeline
from preflight.ingestion.schemas import DocumentType, ExtractorEngine
from preflight.security.rate_limiter import global_rate_limiter


class TestExcelIngestion(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(app, headers={"X-API-Key": "pf_dev_adm_9901"})

    def setUp(self):
        global_rate_limiter.reset()


    def _create_sample_excel_bytes(self) -> bytes:
        """Create a multi-sheet enterprise purchase order Excel file in memory."""
        wb = openpyxl.Workbook()

        # Sheet 1: Cover / Summary (to test automatic skipping)
        ws_cover = wb.active
        ws_cover.title = "Bìa & Hướng dẫn"
        ws_cover["A1"] = "HƯỚNG DẪN ĐẶT HÀNG NHÀ PHÂN PHỐI"
        ws_cover["A2"] = "Vui lòng xem chi tiết tại sheet 'DonDatHang'"

        # Sheet 2: Primary Order Data
        ws_order = wb.create_sheet(title="DonDatHang")

        # Header Info
        ws_order["A1"] = "CÔNG TY CỔ PHẦN BÁN LẺ FPT"
        ws_order["A2"] = "ĐƠN ĐẶT HÀNG (PURCHASE ORDER)"
        ws_order["A4"] = "Số PO:"
        ws_order["B4"] = "PO-2026-XLSX-7788"
        ws_order["D4"] = "Ngày đặt:"
        ws_order["E4"] = "2026-09-01"
        ws_order["A5"] = "Khách hàng:"
        ws_order["B5"] = "FPT Retail Corporation"
        ws_order["D5"] = "Tiền tệ:"
        ws_order["E5"] = "VND"

        # Table Header (Row 7)
        headers = ["STT", "Mã sản phẩm", "Tên hàng hóa", "ĐVT", "Số lượng", "Đơn giá (VND)", "Thành tiền (VND)"]
        for col_idx, h in enumerate(headers, 1):
            ws_order.cell(row=7, column=col_idx, value=h)

        # Line Items
        rows = [
            (1, "LAPTOP-A14", "A14 Business Laptop 14-inch", "Cái", 5, "18.500.000 đ", 92500000),
            (2, "CAB-CAT6-3M", "Dây cáp mạng Cat6 3m bấm sẵn", "Thùng", 10, "1.440.000 đ", 14400000),
            (3, "MONITOR-27", "Màn hình 27-inch 4K UHD", "Cái", 2, "6.200.000 đ", 12400000),
        ]
        for r_offset, r_data in enumerate(rows):
            r_num = 8 + r_offset
            for c_idx, val in enumerate(r_data, 1):
                ws_order.cell(row=r_num, column=c_idx, value=val)

        # Summary Total Row
        ws_order["D11"] = "Tổng cộng:"
        ws_order["G11"] = 119300000

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def test_excel_extractor_parses_multi_sheet_and_line_items(self):
        """Test ExcelExtractor accurately parses sheet, metadata, and all line items."""
        excel_bytes = self._create_sample_excel_bytes()
        extractor = ExcelExtractor()
        extracted, domain_order = extractor.extract(excel_bytes, "PO-FPT-7788.xlsx")

        self.assertEqual(extracted.document_type, DocumentType.SPREADSHEET_EXCEL)
        self.assertEqual(extracted.extractor_used, ExtractorEngine.EXCEL_PARSER)
        self.assertEqual(extracted.header.po_number, "PO-2026-XLSX-7788")
        self.assertEqual(extracted.header.customer, "FPT Retail Corporation")
        self.assertEqual(len(extracted.items), 3)

        # Item 1: Laptop
        self.assertEqual(extracted.items[0].sku, "LAPTOP-A14")
        self.assertEqual(extracted.items[0].quantity, 5)
        self.assertEqual(extracted.items[0].unit_price, Decimal("18500000"))
        self.assertEqual(extracted.items[0].amount, Decimal("92500000"))

        # Item 2: Cable in Carton
        self.assertEqual(extracted.items[1].sku, "CAB-CAT6-3M")
        self.assertEqual(extracted.items[1].quantity, 10)
        self.assertEqual(extracted.items[1].uom, "THÙNG")
        self.assertEqual(extracted.items[1].unit_price, Decimal("1440000"))

        # Math verification
        self.assertTrue(extracted.math_verification.is_valid)
        self.assertEqual(extracted.header.subtotal, Decimal("119300000"))

        # Domain order conversion
        self.assertEqual(domain_order.po_number, "PO-2026-XLSX-7788")
        self.assertEqual(len(domain_order.items), 3)

    def test_intelligent_pipeline_cascading_with_excel(self):
        """Test IntelligentIngestionPipeline end-to-end integration with Excel files."""
        excel_bytes = self._create_sample_excel_bytes()
        pipeline = IntelligentIngestionPipeline()
        extracted, domain_order = pipeline.process_file_bytes(excel_bytes, "order.xlsx")

        self.assertEqual(extracted.document_type, DocumentType.SPREADSHEET_EXCEL)
        self.assertEqual(domain_order.total, Decimal("119300000"))

    def test_api_upload_excel_endpoint(self):
        """Test POST /api/v1/ingest/extract with Excel file."""
        excel_bytes = self._create_sample_excel_bytes()
        files = {
            "file": ("PO-2026-XLSX-7788.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        }
        res = self.client.post("/api/v1/ingest/extract", files=files)
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["document_type"], "SPREADSHEET_EXCEL")
        self.assertEqual(len(data["items"]), 3)
        self.assertTrue(data["math_verification"]["is_valid"])



if __name__ == "__main__":
    unittest.main()
